"""Read-only source ingestion and versioned offline preparation outputs."""

import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import hmac
import importlib.metadata
import json
import os
from pathlib import Path
import secrets
import sys
import time
import uuid

from openpyxl import load_workbook

from .core import make_splits, parse_dialogue


ROOT = Path(__file__).resolve().parents[1]


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def write_jsonl(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='\n') as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, allow_nan=False)+'\n')


def _csv_value(value):
    # JSONL is authoritative. Prefix spreadsheet-executable text only in CSV display files.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'"+value
    return value


def write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key, '')) for key in fields})


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def input_path(source, relative):
    path = (source / relative).resolve()
    if not path.is_relative_to(source) or not path.is_file():
        raise ValueError(f'Input must be an existing file under the source directory: {relative}')
    return path


def read_table(path, sheet, header_row=1, data_row=2):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f'Missing sheet {sheet} in {path.name}')
        ws = workbook[sheet]
        headers = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        if len([x for x in headers if x is not None]) != len(set(x for x in headers if x is not None)):
            raise ValueError(f'Duplicate column names in {path.name}/{sheet}')
        rows = []
        for i, values in enumerate(ws.iter_rows(min_row=data_row, values_only=True), start=data_row):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            row = dict(zip(headers, values))
            row['_row'] = i
            rows.append(row)
        return rows
    finally:
        workbook.close()


def identity(value):
    if value is None or not str(value).strip():
        raise ValueError('Missing student identifier; manual source review required')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def pseudonym(key, term, identifier):
    return 'S-'+hmac.new(key, f'{term}\0{identity(identifier)}'.encode(), hashlib.sha256).hexdigest()[:20]


def stamp(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    result = str(value).strip()
    datetime.fromisoformat(result)  # Malformed dates must fail instead of passing a lexical filter.
    return result


def check_columns(rows, required, name):
    if not rows or not set(required).issubset(rows[0]):
        raise ValueError(f'{name}: required columns missing or no data')


def prepare(project_root=ROOT, config_path=None):
    started = time.perf_counter()
    project_root = Path(project_root).resolve()
    config_path = Path(config_path) if config_path else project_root/'configs/preparation.json'
    config = json.loads(config_path.read_text(encoding='utf-8'))
    source_root = config.get('source_root') or os.environ.get('MATH_HACKATHON_SOURCE_DIR')
    if not isinstance(source_root, str) or not source_root.strip():
        raise ValueError('Set MATH_HACKATHON_SOURCE_DIR or source_root in a local config; no source directory is assumed')
    source = Path(source_root).expanduser().resolve()
    if not source.is_dir():
        raise ValueError('Source directory is unavailable; edit configs/preparation.json')
    if config['model']['remote_calls_enabled'] or config['formal_scoring_enabled']:
        raise ValueError('Preparation does not support network inference or formal scoring')
    start = datetime.fromisoformat(config['spring_start_inclusive']).date()
    end = datetime.fromisoformat(config['spring_end_exclusive']).date()
    if start >= end:
        raise ValueError('Spring date window must be nonempty and ordered')
    local = project_root/'.local'
    local.mkdir(exist_ok=True)
    key_path = local/'pseudonym.key'
    if not key_path.exists():
        if (local/'source_baseline.json').exists() or (project_root/'latest_run.json').exists():
            raise ValueError('Local identity key is missing; recover the original key before continuing')
        key_path.write_bytes(secrets.token_bytes(32))
    key = key_path.read_bytes()
    if len(key) != 32:
        raise ValueError('Invalid local pseudonym key; do not replace it after labels are started')
    manifest = []
    for path in sorted(source.rglob('*')):
        if path.is_file() and not path.name.startswith('~$') and path.suffix.lower() in {'.md', '.xlsx', '.docx'}:
            manifest.append({'path': path.relative_to(source).as_posix(), 'bytes': path.stat().st_size, 'sha256': digest(path)})
    if not manifest:
        raise ValueError('No input documents found')
    source_fingerprints = {x['path']: x['sha256'] for x in manifest}
    baseline_path = local/'source_baseline.json'
    if baseline_path.exists() and json.loads(baseline_path.read_text(encoding='utf-8')) != source_fingerprints:
        raise ValueError('Source files changed since baseline. Review the change and create a new project version; existing labels must not be silently reused.')
    fall_path = input_path(source, config['fall_file'])
    spring_path = input_path(source, config['spring_file'])
    roster_path = input_path(source, config['roster_file'])
    required = ['学号', '问答记录', '问题建立时间', '问答来源']
    fall = read_table(fall_path, '问答记录')
    spring = read_table(spring_path, '问答记录')
    roster = read_table(roster_path, '学生参与情况', 2, 4)
    check_columns(fall, required+['智能体类型'], 'fall')
    check_columns(spring, required, 'spring')
    check_columns(roster, ['学号'], 'roster')
    roster_ids = {identity(x['学号']) for x in roster}
    for row in fall+spring:
        row['_time'] = stamp(row['问题建立时间'])
    usage = read_table(fall_path, '智能体使用次数')
    agents = ['思政点灯人', '探知侠', '数模全才', '数模匹配']
    check_columns(usage, ['学号', '合计']+agents, 'fall usage')
    qa_counts = Counter((identity(x['学号']), x['智能体类型']) for x in fall)
    mismatches = 0
    for row in usage:
        counts = [float(row[agent]) for agent in agents]
        mismatches += int(sum(counts) != float(row['合计']))
        mismatches += sum(count != qa_counts[identity(row['学号']), agent] for count, agent in zip(counts, agents))
    if {identity(x['学号']) for x in usage} != {identity(x['学号']) for x in fall}:
        mismatches += 1
    if mismatches:
        raise ValueError(f'Fall usage and question rows disagree in {mismatches} checks')
    selected_spring = []
    exclusions = Counter()
    for row in spring:
        date = datetime.fromisoformat(row['_time']).date()
        if identity(row['学号']) not in roster_ids:
            exclusions['not_in_primary_roster'] += 1
        elif not start <= date < end:
            exclusions['outside_provisional_window'] += 1
        elif row['问答来源'] != config['source_channel']:
            exclusions['different_channel'] += 1
        else:
            selected_spring.append(row)
    selected_fall = [x for x in fall if x['问答来源'] == config['source_channel']]
    records, turns = [], []
    for term, rows, path in [('fall', selected_fall, fall_path), ('spring', selected_spring, spring_path)]:
        for row in rows:
            record_id = f'{term}-R{row["_row"]:06d}'
            parsed = parse_dialogue(row['问答记录'])
            student_key = pseudonym(key, term, row['学号'])
            records.append({'record_id': record_id, 'term': term, 'student_key': student_key,
                            'source_file': path.relative_to(source).as_posix(), 'sheet': '问答记录',
                            'source_row': row['_row'], 'timestamp': row['_time'],
                            'source_channel': row['问答来源'], 'agent_id': row.get('智能体类型'),
                            'parse_status': parsed['status'], 'parse_reason': parsed['reason'],
                            'candidate_turn_count': len(parsed['turns']), 'semantic_review': 'pending',
                            'raw_text': '' if row['问答记录'] is None else str(row['问答记录'])})
            for i, turn in enumerate(parsed['turns'], 1):
                turns.append(dict(turn, turn_id=f'{record_id}-T{i:03d}', record_id=record_id,
                                  student_key=student_key, term=term, turn_index=i,
                                  timestamp=row['_time'], source_channel=row['问答来源'],
                                  agent_id=row.get('智能体类型'), agent_identity_evidence='source_column' if term == 'fall' else None,
                                  parse_status=parsed['status'], label=None, evidence=None,
                                  confidence=None, review_status='unlabeled'))
    splits = make_splits(turns, config['seed'])
    for row in turns:
        row['pool'] = splits['pool_by_student'][row['student_key']]
        row['component_id'] = splits['component_by_student'][row['student_key']]
    now = datetime.now(timezone.utc)
    run_id = now.strftime('%Y%m%dT%H%M%SZ')+'-'+uuid.uuid4().hex[:6]
    run = project_root/'runs'/run_id
    run.mkdir(parents=True, exist_ok=False)
    write_json(run/'config_snapshot.json', config)
    write_csv(run/'source_manifest.csv', manifest, ['path', 'bytes', 'sha256'])
    write_jsonl(run/'private/records.jsonl', records)
    write_jsonl(run/'private/turns.jsonl', turns)
    write_json(run/'private/sampling_design.json', splits)
    by_id = {x['turn_id']: x for x in turns}
    fields = ['turn_id', 'term', 'record_id', 'turn_index', 'question', 'prior_context',
              'bloom_level', 'evidence_quote', 'outsourcing', 'insufficient_evidence',
              'confidence_1_to_5', 'reviewer_id', 'reviewed_at', 'notes']
    for pool, ids in splits['selected'].items():
        sampled = [by_id[tid] for tid in ids]
        for rater in ['rater_A', 'rater_B']:
            write_csv(run/f'annotations/{pool}_{rater}.csv', sampled, fields)
        write_csv(run/f'annotations/{pool}_adjudication.csv', sampled,
                  fields+['rater_A_label', 'rater_B_label', 'adjudication_reason'])
    review_rows = []
    for term in ['fall', 'spring']:
        candidates = [x for x in records if x['term'] == term and x['parse_status'] not in ['missing', 'url_only']]
        # Deliberately prioritize structural anomalies, then a deterministic random-like sample.
        candidates.sort(key=lambda x: (x['parse_status'] != 'needs_review',
                                      hashlib.sha256((str(config['seed'])+x['record_id']).encode()).hexdigest()))
        review_rows += candidates[:15]
    write_csv(run/'annotations/parse_review.csv', review_rows,
              ['record_id', 'term', 'source_file', 'source_row', 'parse_status', 'parse_reason',
               'candidate_turn_count', 'raw_text', 'roles_correct', 'reviewer_id', 'notes'])
    summary = {
        'state': 'PREPARATION_READY_NOT_EVALUATED', 'run_id': run_id, 'created_utc': now.isoformat(),
        'source_files': len(manifest), 'source_hashes_verified': True, 'fall_usage_mismatches': mismatches,
        'fall': {'raw_records': len(fall), 'unique_ids': len({identity(x['学号']) for x in fall}),
                 'time_min': min(x['_time'] for x in fall), 'time_max': max(x['_time'] for x in fall)},
        'spring_export': {'raw_records': len(spring), 'unique_ids': len({identity(x['学号']) for x in spring}),
                          'time_min': min(x['_time'] for x in spring), 'time_max': max(x['_time'] for x in spring)},
        'spring_roster_size': len(roster_ids), 'spring_exclusion_first_reason': dict(exclusions),
        'population_counts': {'fall': len({identity(x['学号']) for x in usage}), 'spring': len(roster_ids)},
        'selected_records': len(records), 'candidate_turns': len(turns),
        'by_term': {term: {'records': sum(x['term'] == term for x in records),
                          'parse_status': dict(Counter(x['parse_status'] for x in records if x['term'] == term)),
                          'turns': sum(x['term'] == term for x in turns),
                          'students_with_turns': len({x['student_key'] for x in turns if x['term'] == term})}
                    for term in ['fall', 'spring']},
        'sampling_counts': {k: len(v) for k, v in splits['selected'].items()}, 'sampling_strata': splits['strata'],
        'unresolved': ['spring_window_confirmation', 'semantic_role_review', 'near_duplicate_semantic_review',
                       'human_labels', 'annotation_quality', 'agent_identity_spring', 'reference_distribution',
                       'independent_learning_outcomes_and_causal_control'],
        'external_model_calls': 0, 'formal_scores_computed': False,
        'python': sys.version, 'openpyxl': importlib.metadata.version('openpyxl'),
        'elapsed_seconds': round(time.perf_counter()-started, 3),
    }
    summary['runtime_packages'] = {name: importlib.metadata.version(name) for name in ['openpyxl', 'et-xmlfile']}
    write_json(run/'readiness.json', summary)
    code_files = [p for sub in ['challenge', 'tests', 'tools'] for p in (ROOT/sub).rglob('*.py')
                  if '__pycache__' not in p.parts]
    code_files += [p for name in ['run_workbench.ps1', 'run_workbench.sh', 'requirements.txt']
                   if (p := ROOT/name).is_file()]
    write_json(run/'code_manifest.json', [{'path': p.relative_to(ROOT).as_posix(), 'sha256': digest(p)}
                                         for p in sorted(code_files)])
    report = [f'# 准备运行回执：{run_id}', '', '状态：准备完成，尚未进行正式标注或效果评价。', '',
              f'- 源资料：{len(manifest)}个文件，已记录SHA-256；秋季次数表对账差异{mismatches}。',
              f'- 主口径候选：{len(records)}条记录、{len(turns)}个提问轮次。',
              f'- 人工样本：{summary["sampling_counts"]}。',
              '- 每次运行新建目录，不覆盖已填写标注表。身份键保存在工程.local目录，勿删除或公开。',
              '- CSV中的模型/人工标签未填写；不导出学生键或学生级清单，来源清单人数仅按学期汇总。',
              '- 春季日期、角色解析和标注质量仍待确认，当前不能用于正式论文结论。', '',
              '## 立即可做', '',
              '1. 依据docs/标注手册_v1.md填写annotations/parse_review.csv，检查角色与多轮边界。',
              '2. 两名标注者分别填写development_rater_A/B.csv，不互看标签。',
              '3. 第三人用development_adjudication.csv裁决，冻结规则后再打开audit表。',
              '4. 依据docs/workflow.md核实日期、渠道和范围，未核实事项见readiness.json的unresolved；本工程没有代发消息。', '',
              '## 数据范围', '', '```json', json.dumps(summary['by_term'], ensure_ascii=False, indent=2), '```', '']
    (run/'准备回执.md').write_text('\n'.join(report), encoding='utf-8')
    generated = [{'path': p.relative_to(run).as_posix(), 'sha256': digest(p)} for p in sorted(run.rglob('*')) if p.is_file()]
    write_json(run/'generated_manifest.json', generated)
    if not baseline_path.exists():
        write_json(baseline_path, source_fingerprints)
    write_json(project_root/'latest_run.json', {'run_id': run_id, 'path': run.relative_to(project_root).as_posix()})
    return summary


def validate_annotation(path, index_path, expected_ids=None):
    """Validate completed human CSV without filling blanks or accepting invented IDs."""
    index = {}
    with Path(index_path).open(encoding='utf-8') as stream:
        for line in stream:
            item = json.loads(line)
            index[item['turn_id']] = item
    with Path(path).open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream))
    errors, seen, complete = [], set(), 0
    for number, row in enumerate(rows, 2):
        if None in row:
            errors.append(f'row {number}: extra CSV fields; check quoting and commas')
            continue
        row = {key: value if value is not None else '' for key, value in row.items()}
        tid = row.get('turn_id', '')
        if tid in seen or tid not in index:
            errors.append(f'row {number}: duplicate or unknown turn_id')
            continue
        seen.add(tid)
        label = row.get('bloom_level', '').strip()
        if not label:
            errors.append(f'row {number}: label is blank (not completed)')
            continue
        if label not in ['1', '2', '3', '4', '5', '6', 'NA']:
            errors.append(f'row {number}: label must be 1..6 or NA')
            continue
        if not row.get('reviewer_id', '').strip() or not row.get('reviewed_at', '').strip():
            errors.append(f'row {number}: reviewer or date is missing')
            continue
        try:
            datetime.fromisoformat(row['reviewed_at'])
        except ValueError:
            errors.append(f'row {number}: reviewed_at must be an ISO date or timestamp')
            continue
        if row.get('outsourcing') not in ['yes', 'no', 'uncertain'] or row.get('insufficient_evidence') not in ['yes', 'no']:
            errors.append(f'row {number}: outsourcing/insufficient_evidence fields are incomplete')
            continue
        if row.get('confidence_1_to_5') not in ['1', '2', '3', '4', '5']:
            errors.append(f'row {number}: confidence must be 1..5 (subjective, not calibrated)')
            continue
        if (label == 'NA') != (row.get('insufficient_evidence') == 'yes'):
            errors.append(f'row {number}: NA must correspond to insufficient_evidence=yes')
            continue
        evidence = row.get('evidence_quote', '').strip()
        if label != 'NA' and (not evidence or evidence not in index[tid]['question']):
            errors.append(f'row {number}: evidence must be a verbatim span of the current student question')
            continue
        if label == 'NA' and not row.get('notes', '').strip():
            errors.append(f'row {number}: NA requires an explanation')
            continue
        complete += 1
    if not rows:
        errors.append('empty annotation file')
    if expected_ids is not None and seen != set(expected_ids):
        errors.append(f'selected sample mismatch: missing={len(set(expected_ids)-seen)}, extra={len(seen-set(expected_ids))}')
    return {'rows': len(rows), 'completed': complete, 'errors': errors, 'valid': not errors}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    prep = sub.add_parser('prepare')
    prep.add_argument('--config', type=Path)
    validate = sub.add_parser('validate-annotations')
    validate.add_argument('--file', required=True, type=Path)
    validate.add_argument('--index', required=True, type=Path)
    validate.add_argument('--pool', required=True, choices=['development', 'audit', 'risk'])
    args = parser.parse_args()
    try:
        if args.command == 'prepare':
            value = prepare(config_path=args.config)
            print(json.dumps({k: value[k] for k in ['state', 'run_id', 'candidate_turns', 'sampling_counts', 'elapsed_seconds']}, ensure_ascii=False))
        else:
            design = json.loads((args.index.parent/'sampling_design.json').read_text(encoding='utf-8'))
            value = validate_annotation(args.file, args.index, design['selected'][args.pool])
            print(json.dumps(value, ensure_ascii=False))
            return 0 if value['valid'] else 2
    except (OSError, ValueError, KeyError) as exc:
        print(f'Preparation stopped: {exc}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
