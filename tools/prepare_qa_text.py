"""Rebuild the 1,216 question-only inputs from the supplied Excel workbooks."""

from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DEFAULT = Path(os.environ.get("MATH_HACKATHON_SOURCE_DIR", Path.home() / "Desktop" / "黑客松赛题-关老师"))
Q_MARK = re.compile(r"(?m)^Q[:：]\s*")
A_MARK = re.compile(r"(?m)^A[:：]\s*")
TURN_MARK = re.compile(r"(?m)^(Q|A)[:：][ \t]*")
TEST_ROWS = (10, 9, 35, 75, 58, 61, 7, 114, 174, 687)
FIELDS = ("qid", "semester", "agent_or_source", "question_text")


def clean_id(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def extract_questions(transcript: object) -> str | None:
    if not isinstance(transcript, str) or not Q_MARK.search(transcript) or not A_MARK.search(transcript):
        return None
    markers = list(TURN_MARK.finditer(transcript))
    questions = []
    for index, marker in enumerate(markers):
        if marker.group(1) != "Q":
            continue
        end = markers[index + 1].start() if index + 1 < len(markers) else len(transcript)
        question = transcript[marker.end() : end].strip()
        if question:
            questions.append(question)
    return "\n".join(questions) if questions else None


def read_roster(path: Path) -> set[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["学生参与情况"]
        roster = {
            clean_id(row[1])
            for row in sheet.iter_rows(min_row=4, values_only=True)
            if len(row) > 2 and row[2] == "数学模型2026春季班" and clean_id(row[1])
        }
    finally:
        workbook.close()
    if len(roster) != 132:
        raise ValueError(f"2026 roster expected 132 IDs, found {len(roster)}")
    return roster


def read_qa(path: Path, semester: str, roster: set[str] | None = None) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["问答记录"] if "问答记录" in workbook.sheetnames else workbook.active
        header = [str(value) if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
        id_col, text_col = header.index("学号"), header.index("问答记录")
        category_col = header.index("智能体类型" if semester == "2025_fall" else "问答来源")
        records = []
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if roster is not None and clean_id(values[id_col]) not in roster:
                continue
            question_text = extract_questions(values[text_col])
            if not question_text:
                continue
            records.append(
                {
                    "qid": f"{semester}-r{excel_row:06d}",
                    "semester": semester,
                    "agent_or_source": str(values[category_col]).strip(),
                    "question_text": question_text,
                }
            )
    finally:
        workbook.close()
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare the 1,216 question-only Bloom inputs.")
    parser.add_argument("--source-root", type=Path, default=SOURCE_DEFAULT)
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "qa_text.csv")
    args = parser.parse_args()

    source = args.source_root
    path25 = source / "附件2-2025 年秋季学期脱敏数据" / "2025秋-数学模型-问答记录导出20260502导出_已标注_已更新.xlsx"
    folder26 = source / "附件3-2026年春季学期脱敏数据"
    roster = read_roster(folder26 / "请大家对本学期线上线下融合授课-参与数据.xlsx")
    rows25 = read_qa(path25, "2025_fall")
    rows26 = read_qa(folder26 / "问答记录导出0919_已更新.xlsx", "2026_spring", roster)
    if (len(rows25), len(rows26)) != (823, 393):
        raise ValueError(f"Expected 823 + 393 question records, found {len(rows25)} + {len(rows26)}")

    all_rows = rows25 + rows26
    by_qid = {row["qid"]: row for row in all_rows}
    if len(by_qid) != len(all_rows):
        raise ValueError("Duplicate qid in source data")
    first_qids = [f"2025_fall-r{excel_row:06d}" for excel_row in TEST_ROWS]
    ordered = [by_qid[qid] for qid in first_qids] + [row for row in all_rows if row["qid"] not in set(first_qids)]

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(ordered)
    print(f"Prepared {len(rows25)} + {len(rows26)} = {len(ordered)} rows at {args.output}")


if __name__ == "__main__":
    main()
