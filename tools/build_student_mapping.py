"""Resolve qid to existing anonymized student ID in memory only.

The T4 request both proposes a row-level mapping CSV and prohibits any CSV
containing student_key. The stricter privacy rule wins: this module NEVER
writes or prints the mapping. Import ``build_mapping`` from aggregate scripts.
"""

from __future__ import annotations

import csv
import os
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

if __package__:
    from .prepare_qa_text import extract_questions
else:
    from prepare_qa_text import extract_questions


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(os.environ.get("MATH_HACKATHON_SOURCE_DIR", Path.home() / "Desktop" / "黑客松赛题-关老师"))
WORKBOOKS = {
    "2025_fall": SOURCE / "附件2-2025 年秋季学期脱敏数据" / "2025秋-数学模型-问答记录导出20260502导出_已标注_已更新.xlsx",
    "2026_spring": SOURCE / "附件3-2026年春季学期脱敏数据" / "问答记录导出0919_已更新.xlsx",
}
Q_MARK = re.compile(r"(?m)^Q[:：][ \t]*")


@dataclass(frozen=True)
class MappingRecord:
    semester: str
    student_key: str
    n_records: int  # Here: count of source student-Q markers within this qid.


def clean_id(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def qa_index(path: Path = ROOT / "data" / "qa_text.csv") -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or not {"qid", "semester", "question_text"}.issubset(reader.fieldnames):
            raise ValueError("qa_text.csv lacks required mapping checks")
        for row in reader:
            qid, semester = row["qid"].strip(), row["semester"].strip()
            if semester not in WORKBOOKS or not qid.startswith(semester + "-r") or qid in result:
                raise ValueError("Invalid or duplicate qid in qa_text.csv")
            result[qid] = (semester, row["question_text"])
    if len(result) != 1216:
        raise ValueError("Expected 1,216 selected QA records")
    return result


def build_mapping(qids: dict[str, tuple[str, str]] | None = None) -> dict[str, MappingRecord]:
    wanted = qa_index() if qids is None else qids
    result: dict[str, MappingRecord] = {}
    for semester, path in WORKBOOKS.items():
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["问答记录"]
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows))
            id_col, text_col = header.index("学号"), header.index("问答记录")
            for excel_row, values in enumerate(rows, start=2):
                qid = f"{semester}-r{excel_row:06d}"
                if qid not in wanted:
                    continue
                expected_semester, expected_questions = wanted[qid]
                if expected_semester != semester or extract_questions(values[text_col]) != expected_questions:
                    raise ValueError("qid semester or question text differs from source workbook")
                student_key = clean_id(values[id_col])
                question_text = values[text_col]
                if not student_key or not isinstance(question_text, str):
                    raise ValueError("Selected source record lacks student key or transcript")
                result[qid] = MappingRecord(semester, student_key, len(Q_MARK.findall(question_text)))
        finally:
            workbook.close()
    if set(result) != set(wanted):
        raise ValueError(f"Source mapping incomplete: {len(set(wanted) - set(result))} qids missing")
    return result


def main() -> None:
    mapping = build_mapping()
    counts = Counter(row.semester for row in mapping.values())
    students = {semester: {row.student_key for row in mapping.values() if row.semester == semester} for semester in WORKBOOKS}
    print("In-memory anonymized student mapping verified; no mapping CSV written.")
    for semester in WORKBOOKS:
        print(f"{semester}: qids={counts[semester]}, students={len(students[semester])}, "
              f"zero-Q records={sum(row.semester == semester and row.n_records == 0 for row in mapping.values())}")
    print(f"Cross-semester ID overlap={len(students['2025_fall'] & students['2026_spring'])}")


if __name__ == "__main__":
    main()
